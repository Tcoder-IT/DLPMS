import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_path = r"d:\from C\VETHIO\DLPMS\FINALYEAR\Digital_Land_Property_Management_System_Proposal (1) - Formatted.docx"
doc = docx.Document(doc_path)

print("=== VERIFICATION REPORT ===\n")

# 1. Count paragraphs and tables
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"Total sections: {len(doc.sections)}")

# 2. Check margins
print("\n--- Margin Check (first 5 sections) ---")
margin_ok = True
for i, sec in enumerate(doc.sections[:5]):
    top = round(sec.top_margin.inches, 2) if sec.top_margin else None
    bottom = round(sec.bottom_margin.inches, 2) if sec.bottom_margin else None
    left = round(sec.left_margin.inches, 2) if sec.left_margin else None
    right = round(sec.right_margin.inches, 2) if sec.right_margin else None
    ok = (top == 1.0 and bottom == 1.0 and left == 1.0 and right == 1.0)
    if not ok:
        margin_ok = False
    print(f"  Section {i}: top={top}, bottom={bottom}, left={left}, right={right} {'OK' if ok else 'FAIL'}")
print(f"Margins: {'PASS' if margin_ok else 'FAIL'}")

# 3. Check chapter titles
print("\n--- Chapter Title Check ---")
chapter_titles = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text.startswith("CHAPTER ") and (para.style.name == "Title" or "heading" in para.style.name.lower()):
        chapter_titles.append(text)
        runs_bold = all(r.bold for r in para.runs if r.text.strip())
        align = para.alignment
        print(f"  Found: '{text[:60]}' | Bold: {runs_bold} | Align: {align}")
print(f"Total chapter titles found: {len(chapter_titles)}")

# 4. Check for Chapter 5
has_chapter5 = any("CHAPTER FIVE" in p.text for p in doc.paragraphs)
print(f"\nChapter Five exists: {'PASS' if has_chapter5 else 'FAIL'}")

# 5. Check body text formatting (sample 10 body text paragraphs)
print("\n--- Body Text Sample Check ---")
body_issues = 0
body_checked = 0
for i, para in enumerate(doc.paragraphs):
    if para.style.name in ["Body Text", "Normal"] and para.text.strip() and len(para.text.strip()) > 50:
        body_checked += 1
        # Check font
        for run in para.runs:
            if run.text.strip():
                fname = run.font.name
                fsize = run.font.size
                if fname and fname != "Times New Roman":
                    body_issues += 1
                    print(f"  Para {i}: Font '{fname}' (expected Times New Roman)")
                if fsize and round(fsize.pt) != 12:
                    body_issues += 1
                    print(f"  Para {i}: Size {fsize.pt}pt (expected 12pt)")
        if body_checked >= 10:
            break
print(f"Checked {body_checked} body paragraphs, issues: {body_issues}")

# 6. Check heading 1 styling (sample)
print("\n--- Heading 1 (Major Section) Sample Check ---")
h1_checked = 0
h1_issues = 0
for i, para in enumerate(doc.paragraphs):
    if para.style.name == "Heading 1" and para.text.strip():
        h1_checked += 1
        for run in para.runs:
            if run.text.strip():
                fname = run.font.name
                fsize = run.font.size
                is_bold = run.bold
                if fname and fname != "Times New Roman":
                    h1_issues += 1
                if fsize and round(fsize.pt) != 14:
                    h1_issues += 1
                if is_bold is False:
                    h1_issues += 1
        if h1_checked >= 5:
            break
print(f"Checked {h1_checked} Heading 1 paragraphs, issues: {h1_issues}")

# 7. Check no consecutive empty paragraphs
print("\n--- Consecutive Empty Paragraph Check ---")
consecutive_empty = 0
prev_empty = False
for para in doc.paragraphs:
    is_empty = not para.text.strip()
    if is_empty and prev_empty:
        consecutive_empty += 1
    prev_empty = is_empty
print(f"Consecutive empty paragraph occurrences: {consecutive_empty} ({'PASS' if consecutive_empty == 0 else 'FAIL'})")

# 8. Check table formatting (first 3 tables)
print("\n--- Table Check (first 3 tables) ---")
for ti, table in enumerate(doc.tables[:3]):
    print(f"  Table {ti}: {len(table.rows)} rows x {len(table.columns)} cols")
    for row_idx, row in enumerate(table.rows[:2]):
        for col_idx, cell in enumerate(row.cells[:3]):
            for para in cell.paragraphs:
                if para.text.strip():
                    for run in para.runs:
                        if run.text.strip():
                            fname = run.font.name
                            fsize = run.font.size
                            bold = run.bold
                            size_str = f"{round(fsize.pt)}pt" if fsize else "inherited"
                            print(f"    Row {row_idx} Col {col_idx}: font={fname}, size={size_str}, bold={bold}")

print("\n=== VERIFICATION COMPLETE ===")
