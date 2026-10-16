"""
format_docx.py
Academic formatter for the DLPMS Final Year Project Proposal document.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn



def set_cell_background(cell, hex_color):
    """Apply a solid background fill to a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shd element if any
    existing = tcPr.find(qn('w:shd'))
    if existing is not None:
        tcPr.remove(existing)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def has_image(paragraph):
    """Return True if a paragraph contains an inline image/drawing."""
    p_el = paragraph._element
    # Use Clark-notation tags with qn() - no namespaces kwarg needed
    drawing_tag = qn('w:drawing')
    pict_tag    = qn('w:pict')
    for child in p_el.iter():
        if child.tag in (drawing_tag, pict_tag):
            return True
    return False


def set_run_font(run, font_name, font_size_pt, bold=None):
    """Apply font name, size, and bold to a run without touching its text."""
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    if bold is not None:
        run.bold = bold
    # Force font for all character sets (Latin, Complex Script, East Asian)
    rPr = run.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)


def format_runs(paragraph, font_name='Times New Roman', font_size_pt=12, bold=None):
    """Apply formatting to all runs in a paragraph."""
    for run in paragraph.runs:
        set_run_font(run, font_name, font_size_pt, bold)


def apply_paragraph_spacing(para, space_before_pt, space_after_pt, line_spacing=None,
                             first_line_indent=None, left_indent=None, alignment=None):
    """Apply spacing and indentation settings to a paragraph."""
    pf = para.paragraph_format
    pf.space_before = Pt(space_before_pt)
    pf.space_after = Pt(space_after_pt)
    if line_spacing is not None:
        pf.line_spacing = line_spacing
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent
    if left_indent is not None:
        pf.left_indent = left_indent
    if alignment is not None:
        para.alignment = alignment


def insert_chapter_five(doc):
    """Insert CHAPTER FIVE title before the Conclusion paragraph and set heading styles."""
    conclusion_para = None
    conclusion_idx = -1
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip().lower() == 'conclusion' and i > 600:
            conclusion_para = para
            conclusion_idx = i
            break

    if conclusion_para is None:
        print("WARNING: 'Conclusion' paragraph not found. Skipping Chapter Five insertion.")
        return

    print(f"Found Conclusion at paragraph index {conclusion_idx}. Inserting Chapter Five title...")

    # Create a new paragraph element and insert BEFORE the Conclusion paragraph
    conclusion_elem = conclusion_para._element
    new_p_elem = OxmlElement('w:p')
    conclusion_elem.addprevious(new_p_elem)

    # Add a run inside the new paragraph with the chapter title text
    new_r = OxmlElement('w:r')
    new_t = OxmlElement('w:t')
    new_t.text = 'CHAPTER FIVE: CONCLUSION AND RECOMMENDATIONS'
    new_r.append(new_t)
    new_p_elem.append(new_r)

    # Set the style on the new paragraph
    new_pPr = OxmlElement('w:pPr')
    new_pStyle = OxmlElement('w:pStyle')
    new_pStyle.set(qn('w:val'), 'Title')
    new_pPr.insert(0, new_pStyle)
    new_p_elem.insert(0, new_pPr)

    # Set Conclusion and Recommendations to Heading 1
    conclusion_para.style = doc.styles['Heading 1']

    for i in range(conclusion_idx + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip().lower() == 'recommendations':
            print(f"Found Recommendations at paragraph index {i}. Setting to Heading 1...")
            doc.paragraphs[i].style = doc.styles['Heading 1']
            break


def main():
    input_path  = r"d:\from C\VETHIO\DLPMS\FINALYEAR\Digital_Land_Property_Management_System_Proposal (1).docx"
    output_path = r"d:\from C\VETHIO\DLPMS\FINALYEAR\Digital_Land_Property_Management_System_Proposal (1) - Formatted.docx"

    print("Loading document...")
    doc = docx.Document(input_path)

    # ── STEP 1: Insert Chapter Five + fix Conclusion/Recommendations styles ──
    insert_chapter_five(doc)

    # ── STEP 2: Standardise all page margins to 1.0 inch ───────────────────
    print("Standardising page margins...")
    for sec in doc.sections:
        sec.top_margin    = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin   = Inches(1.0)
        sec.right_margin  = Inches(1.0)

    # ── STEP 3: Locate key section boundaries ──────────────────────────────
    toc_start   = -1   # TABLE OF CONTENTS heading
    ack_start   = -1   # ACKNOWLEDGMENT heading
    abbrev_start= -1   # ABBREVIATION AND ACRONYMS heading
    lot_start   = -1   # LIST OF TABLES heading
    lof_start   = -1   # LIST OF FIGURES heading
    intro_start = -1   # First chapter heading (INTRODUCTION / CHAPTER ONE)
    ref_start   = -1   # REFERENCES heading

    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip().upper()
        if 'TABLE OF CONTENTS' in t and toc_start == -1:
            toc_start = i
        elif 'ACKNOWLEDGMENT' in t and ack_start == -1:
            ack_start = i
        elif 'ABBREVIATION AND ACRONYMS' in t and abbrev_start == -1:
            abbrev_start = i
        elif 'LIST OF TABLES' in t and lot_start == -1:
            lot_start = i
        elif 'LIST OF FIGURES' in t and lof_start == -1:
            lof_start = i
        elif t in ('INTRODUCTION', 'CHAPTER ONE: INTRODUCTION') and intro_start == -1:
            intro_start = i
            # Normalise title to CHAPTER ONE: INTRODUCTION
            if t == 'INTRODUCTION':
                # Clear runs and set new text
                for r in para.runs:
                    r.text = ''
                if para.runs:
                    para.runs[0].text = 'CHAPTER ONE: INTRODUCTION'
                else:
                    new_r = OxmlElement('w:r')
                    new_t = OxmlElement('w:t')
                    new_t.text = 'CHAPTER ONE: INTRODUCTION'
                    new_r.append(new_t)
                    para._element.append(new_r)
        elif t == 'REFERENCES' and ref_start == -1:
            ref_start = i

    print(f"Key indices: TOC={toc_start}, ACK={ack_start}, ABBREV={abbrev_start}, "
          f"LOT={lot_start}, LOF={lof_start}, INTRO={intro_start}, REF={ref_start}")

    # ── STEP 4: Format all paragraphs ─────────────────────────────────────
    print("Formatting paragraphs...")
    paragraphs_to_delete = []
    prev_was_heading = False

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        style = para.style.name

        # ─ Cover page (before Table of Contents) ─
        if toc_start != -1 and i < toc_start:
            format_runs(para, font_name='Times New Roman', font_size_pt=12)
            continue

        # ─ Remove extra blank paragraphs (after TOC starts) ─
        if not text and not has_image(para):
            paragraphs_to_delete.append(para)
            continue

        # ─ Detect chapter / front-matter titles ─
        text_upper = text.upper()
        is_chapter_title = (
            text_upper.startswith('CHAPTER ')
            or text_upper in ('TABLE OF CONTENTS', 'ACKNOWLEDGMENT',
                              'ABBREVIATION AND ACRONYMS', 'LIST OF TABLES',
                              'LIST OF FIGURES', 'REFERENCES')
            or style == 'Title'
            or i == intro_start
        )

        if is_chapter_title:
            # Ensure text is stored in uppercase
            if text != text_upper:
                for r in para.runs:
                    r.text = r.text.upper()
            apply_paragraph_spacing(para,
                space_before_pt=24, space_after_pt=18,
                first_line_indent=Inches(0), left_indent=Inches(0),
                alignment=WD_ALIGN_PARAGRAPH.CENTER)
            para.paragraph_format.page_break_before = True
            format_runs(para, font_name='Times New Roman', font_size_pt=16, bold=True)
            prev_was_heading = True
            continue

        # ─ Major section headings (Heading 1) ─
        if style == 'Heading 1':
            apply_paragraph_spacing(para,
                space_before_pt=12, space_after_pt=12,
                first_line_indent=Inches(0), left_indent=Inches(0),
                alignment=WD_ALIGN_PARAGRAPH.LEFT)
            format_runs(para, font_name='Times New Roman', font_size_pt=14, bold=True)
            prev_was_heading = True
            continue

        # ─ Subsection headings (Heading 2 / 3 / 4 / 5) ─
        if style in ('Heading 2', 'Heading 3', 'Heading 4', 'Heading 5'):
            apply_paragraph_spacing(para,
                space_before_pt=6, space_after_pt=6,
                first_line_indent=Inches(0), left_indent=Inches(0),
                alignment=WD_ALIGN_PARAGRAPH.LEFT)
            format_runs(para, font_name='Times New Roman', font_size_pt=12, bold=True)
            prev_was_heading = True
            continue

        # ─ TOC / List of Tables / List of Figures items ─
        in_toc   = toc_start != -1 and ack_start   != -1 and toc_start   < i < ack_start
        in_lot   = lot_start  != -1 and lof_start   != -1 and lot_start   < i < lof_start
        in_lof   = lof_start  != -1 and intro_start != -1 and lof_start   < i < intro_start
        if in_toc or in_lot or in_lof:
            apply_paragraph_spacing(para,
                space_before_pt=2, space_after_pt=2, line_spacing=1.15,
                first_line_indent=Inches(0))
            format_runs(para, font_name='Times New Roman', font_size_pt=12)
            prev_was_heading = False
            continue

        # ─ References section items ─
        if ref_start != -1 and i > ref_start:
            apply_paragraph_spacing(para,
                space_before_pt=0, space_after_pt=6, line_spacing=1.5,
                first_line_indent=Inches(-0.5), left_indent=Inches(0.5),
                alignment=WD_ALIGN_PARAGRAPH.LEFT)
            format_runs(para, font_name='Times New Roman', font_size_pt=12)
            prev_was_heading = False
            continue

        # ─ Paragraphs containing images / figures ─
        if has_image(para):
            apply_paragraph_spacing(para,
                space_before_pt=12, space_after_pt=6,
                first_line_indent=Inches(0), left_indent=Inches(0),
                alignment=WD_ALIGN_PARAGRAPH.CENTER)
            prev_was_heading = False
            continue

        # ─ Table or Figure captions ─
        is_table_caption  = (text.startswith('Table ')  or text.startswith('CRC Card ')
                             or text.startswith('State Chart:'))
        is_figure_caption = text.startswith('Figure ')

        if is_table_caption or is_figure_caption:
            sb = 12 if is_table_caption else 6
            sa = 6  if is_table_caption else 12
            apply_paragraph_spacing(para,
                space_before_pt=sb, space_after_pt=sa,
                first_line_indent=Inches(0), left_indent=Inches(0),
                alignment=WD_ALIGN_PARAGRAPH.CENTER)
            format_runs(para, font_name='Times New Roman', font_size_pt=11, bold=True)
            prev_was_heading = False
            continue

        # ─ Bullet / numbered lists ─
        if style == 'List Paragraph' or text[:1] in ('\u2022', '-', '*'):
            apply_paragraph_spacing(para,
                space_before_pt=0, space_after_pt=4, line_spacing=1.15,
                first_line_indent=Inches(0))
            format_runs(para, font_name='Times New Roman', font_size_pt=12)
            prev_was_heading = False
            continue

        # ─ Regular body text ─
        first_indent = Inches(0) if prev_was_heading else Inches(0.5)
        apply_paragraph_spacing(para,
            space_before_pt=0, space_after_pt=0, line_spacing=1.5,
            first_line_indent=first_indent, left_indent=Inches(0),
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
        format_runs(para, font_name='Times New Roman', font_size_pt=12)
        prev_was_heading = False

    # ── STEP 5: Remove extra blank paragraphs ─────────────────────────────
    print(f"Removing {len(paragraphs_to_delete)} blank paragraphs...")
    for para in paragraphs_to_delete:
        el = para._element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)

    # ── STEP 6: Format tables ──────────────────────────────────────────────
    print("Formatting tables...")
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                # Light grey header row background
                if row_idx == 0:
                    set_cell_background(cell, 'F2F2F2')
                for para in cell.paragraphs:
                    apply_paragraph_spacing(para,
                        space_before_pt=2, space_after_pt=2, line_spacing=1.15,
                        first_line_indent=Inches(0))
                    if row_idx == 0:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        format_runs(para, font_name='Times New Roman',
                                    font_size_pt=11, bold=True)
                    else:
                        # First column: left-align; others: left-align unless very short
                        cell_text = cell.text.strip()
                        if col_idx == 0 or not (cell_text.isdigit() or len(cell_text) <= 5):
                            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        else:
                            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        format_runs(para, font_name='Times New Roman',
                                    font_size_pt=11, bold=False)

    # ── STEP 7: Fix figure caption placement (ensure caption is BELOW figure) ──
    print("Re-checking caption placement...")
    body = doc.element.body
    elements = list(body)
    i = 0
    while i < len(elements) - 1:
        el = elements[i]
        # Look for a paragraph with an image
        if el.tag == qn('w:p'):
            para = docx.text.paragraph.Paragraph(el, doc)
            if has_image(para) and i > 0:
                prev_el = elements[i - 1]
                if prev_el.tag == qn('w:p'):
                    prev_para = docx.text.paragraph.Paragraph(prev_el, doc)
                    pt = prev_para.text.strip()
                    if pt.startswith('Figure '):
                        print(f"  Moving figure caption from above to below: '{pt[:50]}'")
                        body.remove(prev_el)
                        body.insert(i, prev_el)
                        elements = list(body)
                        continue
        i += 1

    # Fix table captions that appear BELOW a table (move them above)
    elements = list(body)
    i = len(elements) - 1
    while i > 0:
        el      = elements[i]
        prev_el = elements[i - 1]
        if el.tag == qn('w:p') and prev_el.tag == qn('w:tbl'):
            para = docx.text.paragraph.Paragraph(el, doc)
            t = para.text.strip()
            if t.startswith('Table ') or t.startswith('CRC Card ') or t.startswith('State Chart:'):
                print(f"  Moving table caption from below to above: '{t[:50]}'")
                body.remove(el)
                body.insert(i - 1, el)
                elements = list(body)
                i = len(elements) - 1
                continue
        i -= 1

    # ── STEP 8: Save ───────────────────────────────────────────────────────
    print("Saving formatted document...")
    doc.save(output_path)
    print(f"Done! Saved to:\n  {output_path}")


if __name__ == '__main__':
    main()
